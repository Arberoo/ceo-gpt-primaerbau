#!/usr/bin/env python3
"""
Kontoauszug abhaken (PrimaerBau, 17.10.2026).

Fuer jeden Monat jede Bankbuchung mit Haken: Beleg liegt in DATEV, braucht
keinen Beleg, oder fehlt. Grundlage:
  - abschluss/JJJJ/MM Monat/DATEV-Abgleich_JJJJ-MM.csv  (aus datev_abgleich.py)
  - zusaetzlich hochgeladene Belege seit dem letzten DATEV-Lesen:
    CSV mit Spalten bank_buchung, bank_betrag, datei (z. B. bons_abgleich.csv)
Ausgabe: Excel-Datei mit einem Blatt je Monat, farbig.

    .venv/Scripts/python.exe -X utf8 .claude/skills/buchhaltung/scripts/kontoauszug_abhaken.py <ziel.xlsx> 2026-09 2026-10 --nachgereicht <csv> [...]
"""
import csv
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
CFG = json.load(open(os.path.join(os.path.dirname(HERE), "buchhaltung.config.json"), encoding="utf-8"))
MONATE = {"01": "01 Januar", "02": "02 Februar", "03": "03 Maerz", "04": "04 April", "05": "05 Mai",
          "06": "06 Juni", "07": "07 Juli", "08": "08 August", "09": "09 September",
          "10": "10 Oktober", "11": "11 November", "12": "12 Dezember"}
GRUEN, GRAU, ROT, GELB = "C6EFCE", "EDEDED", "FFC7CE", "FFEB9C"


def nachgereicht_laden(pfade):
    n = {}
    for p in pfade:
        for r in csv.DictReader(open(p, encoding="utf-8")):
            if r.get("bank_buchung"):
                # optional Spalte status: "fehlt" = von Hand geprueft, Beleg fehlt in DATEV
                wert = r.get("datei") or r.get("bild") or ""
                if (r.get("status") or "").strip() == "fehlt":
                    wert = "FEHLT"
                n.setdefault((r["bank_buchung"], f"{float(r['bank_betrag']):.2f}"), []).append(wert)
    return n


def main():
    args = sys.argv[1:]
    ziel = args[0]
    nach = []
    if "--nachgereicht" in args:
        i = args.index("--nachgereicht")
        nach, args = args[i + 1:], args[:i]
    monate = args[1:]
    extra = nachgereicht_laden(nach)
    wb = Workbook()
    wb.remove(wb.active)
    gesamt = {}
    for m in monate:
        jahr, mm = m.split("-")
        datei = os.path.join(CFG["folder_root"], jahr, MONATE[mm], f"DATEV-Abgleich_{m}.csv")
        rows = list(csv.DictReader(open(datei, encoding="utf-8")))
        ws = wb.create_sheet(MONATE[mm])
        ws.append(["✓", "Datum", "Betrag €", "Gegenpartei / Verwendungszweck", "Status", "Beleg / Datei", "Kategorie (Vorschlag)"])
        for c in ws[1]:
            c.font = Font(bold=True)
        z = {"in DATEV": 0, "belegfrei": 0, "fehlt": 0, "nur lokal": 0}
        for r in sorted(rows, key=lambda x: x["Datum"]):
            key = (r["Datum"], f"{float(r['Betrag']):.2f}")
            if r["Beleg_Status"] == "kein Beleg noetig":
                haken, status, beleg, farbe = "✓", "braucht keinen Beleg", "", GRAU
                z["belegfrei"] += 1
            elif key in extra and extra[key] and extra[key][0] == "FEHLT":
                extra[key].pop(0)
                haken, status, beleg, farbe = "✗", "Beleg fehlt (von Hand geprueft)", "", ROT
                z["fehlt"] += 1
            elif key in extra and extra[key]:
                # von Hand geprueft (selbst hochgeladen oder in DATEV gefunden): Vorrang vor dem Betrags-Abgleich
                haken, status, beleg, farbe = "✓", "Beleg in DATEV (von Hand geprueft)", extra[key].pop(0), GRUEN
                z["in DATEV"] += 1
            elif r["DATEV"] not in ("-", "NICHT IN DATEV"):
                haken, status, beleg, farbe = "✓", "Beleg in DATEV", r["DATEV_Datei"], GRUEN
                z["in DATEV"] += 1
            elif r["Beleg_Status"] == "Beleg vorhanden":
                haken, status, beleg, farbe = "", "Beleg liegt bei uns, in DATEV nicht sicher gefunden - bitte pruefen", r["Beleg_Datei"], GELB
                z["nur lokal"] += 1
            else:
                haken, status, beleg, farbe = "✗", "Beleg fehlt", "", ROT
                z["fehlt"] += 1
            ws.append([haken, r["Datum"], float(r["Betrag"]), r["Gegenpartei"][:120], status, beleg, r["Kategorie_Vorschlag"]])
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=farbe)
        for spalte, breite in zip("ABCDEFG", (4, 11, 12, 70, 34, 55, 28)):
            ws.column_dimensions[spalte].width = breite
        for c in ws["C"][1:]:
            c.number_format = '#,##0.00'
        ws.freeze_panes = "A2"
        ws.append([])
        ws.append(["", "Summe", "", f"{len(rows)} Buchungen: {z['in DATEV']} mit Beleg in DATEV, "
                   f"{z['belegfrei']} brauchen keinen, {z['nur lokal']} nur bei uns, {z['fehlt']} fehlen"])
        ws.cell(ws.max_row, 4).alignment = Alignment(wrap_text=True)
        gesamt[m] = z
        print(m, z)
    wb.save(ziel)
    print(ziel)


if __name__ == "__main__":
    main()
