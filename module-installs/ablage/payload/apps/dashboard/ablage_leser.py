"""Leser der automatischen Ablage: aus jeder Datei Text und Metadaten ziehen.

REINE MECHANIK, KEIN URTEIL. Dieser Teil ist bei jedem Kunden identisch und
deshalb sauber vom Rest getrennt: was eine Datei ENTHAELT, ist eine Tatsache;
wohin sie GEHOERT, ist ein Urteil und lebt in ablage.py.

Wege, billig vor teuer:
    PDF mit Textebene -> fitz (PyMuPDF), kostet nichts
    PDF ohne Textebene -> erste Seiten als Bild rendern -> Vision (kostet)
    DOCX / XLSX -> python-docx / openpyxl
    Bilder -> Vision
    TXT / MD / CSV -> direkt

Zwei Dinge, die nicht wegoptimiert werden duerfen:

1. JEDE KUERZUNG WIRD MARKIERT. Ein Modell, das nicht weiss, dass es nur die
   ersten 20 von 120 Seiten sieht, urteilt selbstbewusst falsch. Dieselbe
   Lehre wie bei belege_suchen (24.08.2026), wo ein Werkzeug nicht sagte,
   wonach es gesucht hatte, und der Agent das leere Ergebnis fuer die
   Antwort hielt.
2. JEDER ZWEIG FAENGT SEINE EIGENEN FEHLER AB. Eine unlesbare Datei darf den
   ganzen Stapel nie kippen. Sie kommt mit leerem Text und gesetztem `fehler`
   zurueck und landet dadurch im Auffang - sichtbar, nicht verschwunden.
"""
from datetime import datetime
from pathlib import Path

MAX_ZEICHEN = 20000 # reicht fuer die Einordnung; ganze Vertraege braucht sie nicht
MAX_SEITEN = 20 # PDF-Seiten fuer die Textebene
VISION_SEITEN = 2 # gescannte PDFs: so viele Seiten durch die Bild-Erkennung
MIN_TEXT = 100 # darunter gilt ein PDF als gescannt (keine Textebene)


def _kuerzen(text, grenze=MAX_ZEICHEN, was="Zeichen"):
    text = (text or "").strip()
    if len(text) <= grenze:
        return text
    return text[:grenze] + f"\n\n… [gekürzt: {grenze} von {len(text)} {was}]"


def _pdf(pfad):
    import fitz
    meta, teile, wie = {}, [], "text"
    with fitz.open(str(pfad)) as doc:
        roh = doc.metadata or {}
        for k in ("title", "author", "subject", "creationDate", "modDate"):
            if roh.get(k):
                meta[k] = str(roh[k])[:200]
        meta["seiten"] = doc.page_count
        for i, page in enumerate(doc):
            if i >= MAX_SEITEN:
                teile.append(f"\n… [gekürzt: {MAX_SEITEN} von {doc.page_count} Seiten gelesen]")
                break
            teile.append(page.get_text())
        text = "\n".join(teile).strip()

        # Gescannt: keine oder fast keine Textebene -> Bild-Erkennung.
        if len(text) < MIN_TEXT and doc.page_count:
            bilder = []
            for i in range(min(VISION_SEITEN, doc.page_count)):
                pix = doc[i].get_pixmap(dpi=150)
                bilder.append(pix.tobytes("png"))
            erkannt = [_vision(b) for b in bilder]
            erkannt = [e for e in erkannt if e]
            if erkannt:
                text = "\n\n".join(erkannt)
                wie = "vision"
                if doc.page_count > VISION_SEITEN:
                    text += (f"\n\n… [gescanntes PDF: {VISION_SEITEN} von "
                             f"{doc.page_count} Seiten erkannt]")
    return _kuerzen(text), meta, wie


def _docx(pfad):
    import docx
    d = docx.Document(str(pfad))
    teile = [p.text for p in d.paragraphs if p.text.strip()]
    for tab in d.tables:
        for row in tab.rows:
            zellen = [c.text.strip() for c in row.cells if c.text.strip()]
            if zellen:
                teile.append(" | ".join(zellen))
    meta = {}
    try:
        cp = d.core_properties
        for k in ("title", "author", "subject"):
            if getattr(cp, k, None):
                meta[k] = str(getattr(cp, k))[:200]
        if cp.created:
            meta["created"] = cp.created.isoformat()
    except Exception:
        pass
    return _kuerzen("\n".join(teile)), meta, "text"


def _xlsx(pfad):
    import openpyxl
    wb = openpyxl.load_workbook(str(pfad), read_only=True, data_only=True)
    teile = []
    try:
        for ws in wb.worksheets:
            teile.append(f"### Blatt: {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 40:
                    teile.append("… [weitere Zeilen ausgelassen]")
                    break
                werte = [str(v) for v in row if v is not None]
                if werte:
                    teile.append(" | ".join(werte))
    finally:
        wb.close()
    return _kuerzen("\n".join(teile)), {"blaetter": len(wb.sheetnames)}, "text"


def _text(pfad):
    roh = Path(pfad).read_text(encoding="utf-8", errors="replace")
    return _kuerzen(roh), {}, "text"


_VISION_PROMPT = (
    "Gib den INHALT dieses Dokuments wieder. Beginne mit der Dokumentart und dem "
    "Titel, dann alle erkennbaren Kerndaten (Namen, Firmen, Nummern, Beträge, "
    "Daten, Fristen) und eine kurze Zusammenfassung. Gib Text, der im Bild steht, "
    "WÖRTLICH wieder. Erfinde nichts, was nicht zu sehen ist; ist etwas unleserlich, "
    "schreib das dazu."
)


def _vision(bild):
    """Bild (bytes oder Pfad) durch die Bild-Erkennung. Leer bei Fehler.

    claude_llm.complete nimmt Pfad ODER bytes, KEIN base64 (ein base64-String
    landet dort als Dateiname: 'File name too long').
    """
    try:
        import sys
        ROOT = Path(__file__).resolve().parents[2]
        sys.path.insert(0, str(ROOT / "scripts"))
        from claude_llm import complete
        return (complete(_VISION_PROMPT, image=bild, agent="ablage_leser") or "").strip()
    except Exception:
        return ""


def _bild(pfad):
    daten = Path(pfad).read_bytes()
    try:
        from apps.dashboard.aufgaben_maschine import _resize_for_vision
        daten = _resize_for_vision(daten)
    except Exception:
        pass # Original schicken ist besser als gar nichts
    text = _vision(daten)
    return _kuerzen(text), {}, "vision"


_WEGE = {
    ".pdf": _pdf, ".docx": _docx, ".doc": _docx,
    ".xlsx": _xlsx, ".xls": _xlsx,
    ".txt": _text, ".md": _text, ".csv": _text,
}


def lies(pfad, orig_name=None):
    """-> {"text", "meta", "wie", "fehler"}. Wirft nie."""
    p = Path(pfad)
    erg = {"text": "", "meta": {}, "wie": "keiner", "fehler": None}
    try:
        st = p.stat()
        erg["meta"] = {
            "dateiname": orig_name or p.name,
            "groesse_kb": round(st.st_size / 1024, 1),
            "geaendert": datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
        }
    except Exception as exc:
        erg["fehler"] = f"Datei nicht lesbar: {exc}"
        return erg

    suffix = p.suffix.lower()
    fn = _WEGE.get(suffix)
    try:
        if fn:
            text, meta, wie = fn(p)
        elif suffix in (".png", ".jpg", ".jpeg", ".heic", ".webp"):
            text, meta, wie = _bild(p)
        else:
            erg["fehler"] = f"kein Leser für {suffix}"
            return erg
        erg["text"] = text
        erg["meta"].update(meta)
        erg["wie"] = wie
        if not text.strip():
            erg["fehler"] = "kein Text erkannt"
    except Exception as exc:
        erg["fehler"] = f"{type(exc).__name__}: {exc}"
    return erg
